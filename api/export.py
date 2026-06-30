"""
Export API
- GET /api/export/<session_id>/csv   → CSV download
- GET /api/export/<session_id>/excel → Excel (.xlsx) download
"""
import csv
import io
import os

from flask import Blueprint, Response, send_file
from database import get_db

bp = Blueprint("export", __name__)

_COLS = [
    "id", "latitude", "longitude", "status_tanam",
    "diameter_tajuk_m", "kategori_tajuk", "usia_bulan", "blok_id",
]

_QUERY = (
    "SELECT l.id, l.latitude, l.longitude, l.status_tanam, "
    "       l.diameter_tajuk_m, l.kategori_tajuk, l.usia_bulan, l.blok_id "
    "FROM lubang_deteksi l "
    "WHERE l.session_id=? ORDER BY l.id"
)


# ---------------------------------------------------------------------------
# CSV export
# ---------------------------------------------------------------------------

@bp.route("/<int:sid>/csv", methods=["GET"])
def export_csv(sid):
    db = get_db()
    rows = db.execute(_QUERY, (sid,)).fetchall()
    sesi = db.execute(
        "SELECT nama, tanggal_terbang FROM drone_sessions WHERE id=?", (sid,)
    ).fetchone()
    db.close()

    output = _build_csv(rows)

    tanggal = sesi["tanggal_terbang"] if sesi else "export"
    fname = f"lubang_{tanggal}.csv"
    return Response(
        output,
        mimetype="text/csv",
        headers={"Content-Disposition": f"attachment; filename={fname}"},
    )


def _build_csv(rows) -> str:
    """Buat CSV string dari list row (dict-like). Dapat diuji tanpa Flask."""
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(_COLS)
    for r in rows:
        writer.writerow([
            r["id"],
            r["latitude"],
            r["longitude"],
            r["status_tanam"],
            r["diameter_tajuk_m"] if r["diameter_tajuk_m"] is not None else "",
            r["kategori_tajuk"],
            r["usia_bulan"] if r["usia_bulan"] is not None else "",
            r["blok_id"] if r["blok_id"] is not None else "",
        ])
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

@bp.route("/<int:sid>/excel", methods=["GET"])
def export_excel(sid):
    db = get_db()
    rows = db.execute(_QUERY, (sid,)).fetchall()
    sesi = db.execute(
        "SELECT nama, tanggal_terbang FROM drone_sessions WHERE id=?", (sid,)
    ).fetchone()
    db.close()

    xlsx_bytes = _build_excel(rows)

    tanggal = sesi["tanggal_terbang"] if sesi else "export"
    fname = f"lubang_{tanggal}.xlsx"

    buf = io.BytesIO(xlsx_bytes)
    buf.seek(0)
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=fname,
    )


def _build_excel(rows) -> bytes:
    """Buat Excel bytes dari list row (dict-like). Dapat diuji tanpa Flask."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lubang Deteksi"

    # Header
    headers = _COLS
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Data
    for ri, r in enumerate(rows, 2):
        ws.cell(ri, 1, r["id"])
        ws.cell(ri, 2, r["latitude"])
        ws.cell(ri, 3, r["longitude"])
        ws.cell(ri, 4, r["status_tanam"])
        ws.cell(ri, 5, r["diameter_tajuk_m"])
        ws.cell(ri, 6, r["kategori_tajuk"])
        ws.cell(ri, 7, r["usia_bulan"])
        ws.cell(ri, 8, r["blok_id"])

    # Auto column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                val_len = len(str(cell.value)) if cell.value is not None else 0
                max_len = max(max_len, val_len)
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = max(max_len + 2, 10)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
